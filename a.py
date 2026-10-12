from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Ejercicio de investigación de Python")
run.bold = True
run.font.size = Pt(16)

sections = [
("1. Defina una lista y una tupla en Python. Proporcione algunos ejemplos.",
"""Una lista (list) es una colección ordenada y mutable de elementos. Esto significa que sus elementos pueden modificarse, agregarse o eliminarse después de crearla.

Ejemplos:
frutas = ["manzana", "pera", "naranja"]
numeros = [1, 2, 3, 4, 5]

frutas.append("plátano")
frutas[0] = "uva"

Una tupla (tuple) también es una colección ordenada, pero es inmutable, por lo que sus elementos no pueden modificarse después de crearla.

Ejemplos:
coordenadas = (10, 20)
colores = ("rojo", "verde", "azul")
dias = ("lunes", "martes", "miércoles")

La principal diferencia es que una lista puede modificarse, mientras que una tupla no."""),
("2. ¿Qué es un espacio de nombres en Python?",
"""Un espacio de nombres (namespace) es un sistema que Python utiliza para organizar y distinguir los nombres de variables, funciones, clases y otros objetos.

Por ejemplo:
nombre = "Martín"

def saludar():
    mensaje = "Hola"

En este caso, nombre y mensaje son nombres asociados a diferentes objetos. Python utiliza espacios de nombres para evitar conflictos entre nombres iguales que pertenecen a diferentes contextos.

Existen, entre otros:
• Espacio de nombres local: pertenece a una función.
• Espacio de nombres global: pertenece al módulo o programa.
• Espacio de nombres incorporado (built-in): contiene elementos propios de Python, como print(), len() e int()."""),
("3. ¿Cuál es la diferencia entre una variable local y una variable global?",
"""Una variable local es una variable creada dentro de una función y normalmente solo puede utilizarse dentro de ella.

Ejemplo:
def calcular():
    numero = 10
    print(numero)

calcular()

En este ejemplo, numero es una variable local.

Una variable global se define fuera de las funciones y puede ser utilizada desde diferentes partes del programa.

Ejemplo:
numero = 10

def mostrar():
    print(numero)

mostrar()

Aquí numero es una variable global.

Diferencia principal:
• Variable local: se define dentro de una función y su alcance está limitado principalmente a esa función.
• Variable global: se define fuera de las funciones y pertenece al espacio de nombres global del módulo."""),
("4. ¿Qué es un IDE? Mencione algunos IDE comunes que podrían utilizarse con Python.",
"""Un IDE (Integrated Development Environment), o Entorno de Desarrollo Integrado, es un programa que proporciona herramientas para desarrollar software en un solo entorno.

Normalmente incluye:
• Editor de código.
• Autocompletado.
• Depurador (debugger).
• Ejecución del programa.
• Gestión de proyectos.
• Detección de errores y resaltado de sintaxis.

Algunos IDE o entornos utilizados para Python son:
• PyCharm
• Visual Studio Code
• Spyder
• IDLE, incluido con Python.
• JupyterLab, muy utilizado para análisis de datos y programación interactiva."""),
("5. ¿Qué son los módulos en Python? Proporcione algunos ejemplos.",
"""Un módulo es un archivo que contiene código Python que puede ser utilizado desde otro programa. Los módulos permiten organizar y reutilizar código.

Por ejemplo:
import math

print(math.sqrt(25))

Aquí se importa el módulo math, que proporciona funciones matemáticas.

Otros ejemplos de módulos estándar son:
• math → operaciones matemáticas.
• random → generación de valores aleatorios.
• json → trabajar con datos JSON.
• os → interacción con el sistema operativo.
• datetime → trabajar con fechas y horas.

También podemos crear nuestros propios módulos. Por ejemplo, un archivo llamado calculadora.py podría contener:

def sumar(a, b):
    return a + b

Y utilizarlo desde otro archivo:

import calculadora

resultado = calculadora.sumar(5, 3)
print(resultado)"""),
("6. ¿Cuál es la diferencia entre una matriz y una lista?",
"""Una lista es una estructura de datos que permite almacenar una colección de elementos y puede contener elementos de diferentes tipos.

Ejemplo:
lista = [1, "Python", 3.14, True]

Una matriz representa normalmente datos organizados en filas y columnas, es decir, en dos o más dimensiones. En Python, las matrices pueden representarse mediante listas anidadas:

matriz = [
    [1, 2, 3],
    [4, 5, 6],
    [7, 8, 9]
]

La matriz anterior tiene 3 filas y 3 columnas.

Para trabajar con matrices y operaciones matemáticas de forma más avanzada también se utiliza la biblioteca NumPy:

import numpy as np

matriz = np.array([
    [1, 2],
    [3, 4]
])

Diferencia principal: una lista es una colección general de elementos, mientras que una matriz está estructurada principalmente para representar datos multidimensionales organizados en filas y columnas."""),
("7. ¿Qué son los operadores? Proporcione algunos ejemplos.",
"""Los operadores son símbolos o palabras que permiten realizar operaciones sobre valores y variables.

Operadores aritméticos:
a = 10
b = 3

a + b   # Suma
a - b   # Resta
a * b   # Multiplicación
a / b   # División
a % b   # Módulo o resto
a ** b  # Potencia

Operadores de comparación:
a == b  # Igual
a != b  # Diferente
a > b   # Mayor que
a < b   # Menor que
a >= b  # Mayor o igual
a <= b  # Menor o igual

Operadores lógicos:
x = True
y = False

x and y
x or y
not x

Operadores de asignación:
x = 10
x += 5   # x = x + 5
x -= 2   # x = x - 2
x *= 3   # x = x * 3

En resumen, los operadores permiten realizar cálculos, comparaciones, asignaciones y operaciones lógicas, siendo fundamentales para construir la lógica de un programa en Python.""")
]

for heading, body in sections:
    p = doc.add_paragraph()
    r = p.add_run(heading)
    r.bold = True
    r.font.size = Pt(12)
    for line in body.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(line)
        if line.startswith(("frutas =", "numeros =", "coordenadas =", "colores =", "dias =", "def ", "    ", "nombre =", "numero =", "print(", "import ", "resultado =", "lista =", "matriz =", "a = ", "b = ", "x = ", "y = ", "a +", "a -", "a *", "a /", "a %", "a **", "a ==", "a !=", "a >", "a <", "a >=", "a <=", "x and", "x or", "not x")):
            r.font.name = "Courier New"
            r.font.size = Pt(10)

path = "/mnt/data/Ejercicio_investigacion_Python.docx"
doc.save(path)
print(path)
